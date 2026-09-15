import os
import re
import uuid
import textwrap
from typing import Dict, Any, List, Tuple, Optional
from django.conf import settings
import docx
from docx.shared import RGBColor

class SmartCodeFormatter:
    """
    AI-Heuristic Code Indenter & Formatter
    Preserves existing indentations, dedents common margins, repairs first-line-stripped anomalies,
    or automatically reconstructs syntax-correct indentation for Python, C++, SQL, and HTML when input code is flat.
    """

    @classmethod
    def clean_and_dedent(cls, code: str) -> str:
        """
        Normalizes line breaks, tabs, non-breaking spaces, strips leading/trailing blank lines,
        and removes common leading indentation across all non-blank lines.
        """
        if not code or not code.strip():
            return ''
        normalized = code.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ').replace('\xa0', ' ')
        raw_lines = normalized.split('\n')
        while raw_lines and not raw_lines[0].strip():
            raw_lines.pop(0)
        while raw_lines and not raw_lines[-1].strip():
            raw_lines.pop()
        if not raw_lines:
            return ''
        return textwrap.dedent('\n'.join(raw_lines))

    @classmethod
    def repair_first_line_stripped(cls, code: str, lang: str = 'python') -> str:
        """
        Detects and fixes anomalies where line 1 has 0 indentation (e.g. from an unintended trim() or Word margin),
        while line 1 did NOT open a block, and lines 2+ all share an extra common indentation (e.g. 4 spaces).
        """
        lines = [l.rstrip() for l in code.split('\n')]
        non_blank_indices = [i for i, l in enumerate(lines) if l.strip()]
        if len(non_blank_indices) < 2:
            return code

        first_idx = non_blank_indices[0]
        first_line = lines[first_idx]
        first_indent = len(first_line) - len(first_line.lstrip())
        first_stripped = first_line.strip()

        opens_block = False
        lang_lower = (lang or 'python').lower()
        if lang_lower in ['python', 'py', 'python3']:
            if first_stripped.endswith(':') or re.match(r'^(def\b|class\b|if\b|for\b|while\b|try\b|with\b|async\b)', first_stripped):
                opens_block = True
        elif lang_lower in ['cpp', 'c++', 'c', 'javascript', 'js', 'java', 'css']:
            if first_stripped.endswith('{') or '{' in first_stripped:
                opens_block = True

        if not opens_block and first_indent == 0:
            subsequent_indices = non_blank_indices[1:]
            subsequent_indents = [len(lines[i]) - len(lines[i].lstrip()) for i in subsequent_indices]
            min_subsequent_indent = min(subsequent_indents)

            if min_subsequent_indent >= 2:
                repaired_lines = list(lines)
                for i in subsequent_indices:
                    if repaired_lines[i].strip():
                        repaired_lines[i] = lines[i][min_subsequent_indent:]
                return '\n'.join(repaired_lines)

        return code

    @classmethod
    def auto_indent_python(cls, code: str) -> str:
        code = cls.clean_and_dedent(code)
        code = cls.repair_first_line_stripped(code, 'python')

        lines = [l.rstrip() for l in code.split('\n')]
        non_blank = [l for l in lines if l.strip()]
        if not non_blank:
            return code

        # If already has consistent indentation, return dedented and repaired lines
        indented_count = sum(1 for l in non_blank if l.startswith('  '))
        if indented_count >= max(1, len(non_blank) // 3):
            return '\n'.join(lines)

        indent_level = 0
        formatted_lines = []
        stack = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                formatted_lines.append('')
                continue

            # 1. Dedent keywords (elif, else, except, finally)
            if re.match(r'^(elif\b|else\s*:|except\b|finally\s*:)', stripped):
                if stack:
                    if stripped.startswith('elif') or stripped.startswith('else'):
                        while stack and stack[-1]['type'] not in ['if', 'elif', 'for', 'while', 'try']:
                            stack.pop()
                        if stack:
                            indent_level = stack[-1]['indent']
                    elif stripped.startswith('except') or stripped.startswith('finally'):
                        while stack and stack[-1]['type'] not in ['try', 'except']:
                            stack.pop()
                        if stack:
                            indent_level = stack[-1]['indent']
            else:
                # 2. Previous line was return / break / continue
                if formatted_lines:
                    prev_line = formatted_lines[-1].strip()
                    if prev_line.startswith('return ') or prev_line == 'return' or prev_line == 'break' or prev_line == 'continue':
                        if stack and stack[-1]['type'] in ['if', 'elif', 'else']:
                            stack.pop()
                            indent_level = (stack[-1]['indent'] + 1) if stack else 0

                # 3. Previous statement closed while loop with loop var update (e.g. j -= 1)
                if stack and stack[-1]['type'] == 'while':
                    prev_line = formatted_lines[-1].strip() if formatted_lines else ''
                    loop_var = stack[-1].get('loop_var')
                    if loop_var and re.search(rf'\b{loop_var}\s*(-=|\+=|=)\s*', prev_line):
                        if not re.search(rf'\b{loop_var}\s*(-=|\+=|=)\s*', stripped):
                            stack.pop()
                            indent_level = stack[-1]['indent'] if stack else 0

                # 4. Previous statement closed for/while loop (e.g. print(...) or return after loop)
                if stack and stack[-1]['type'] in ['for', 'while']:
                    if stripped.startswith('return ') or stripped.startswith('print('):
                        stack.pop()
                        indent_level = stack[-1]['indent'] if stack else 0

            formatted_lines.append(('    ' * indent_level) + stripped)

            # If line ends with ':', opens a new block
            if stripped.endswith(':') or re.search(r':\s*(?:#.*)?$', stripped):
                block_type = 'block'
                loop_var = ''
                if stripped.startswith('def '): block_type = 'def'
                elif stripped.startswith('class '): block_type = 'class'
                elif stripped.startswith('for '):
                    block_type = 'for'
                    m = re.match(r'for\s+([a-zA-Z_]\w*)\s+in', stripped)
                    if m: loop_var = m.group(1)
                elif stripped.startswith('while '):
                    block_type = 'while'
                    m = re.match(r'while\s+([a-zA-Z_]\w*)', stripped)
                    if m: loop_var = m.group(1)
                elif stripped.startswith('if '): block_type = 'if'
                elif stripped.startswith('elif '): block_type = 'elif'
                elif stripped.startswith('else'): block_type = 'else'
                elif stripped.startswith('try'): block_type = 'try'
                elif stripped.startswith('except'): block_type = 'except'
                elif stripped.startswith('finally'): block_type = 'finally'

                stack.append({'type': block_type, 'indent': indent_level, 'loop_var': loop_var})
                indent_level += 1

        return '\n'.join(formatted_lines)

    @classmethod
    def auto_indent_cpp(cls, code: str) -> str:
        lines = [l.rstrip() for l in code.split('\n')]
        non_blank = [l for l in lines if l.strip()]
        if not non_blank:
            return code

        indented_count = sum(1 for l in non_blank if l.startswith('  ') or l.startswith('\t'))
        if indented_count >= max(1, len(non_blank) // 3):
            return '\n'.join(l.replace('\t', '    ') for l in lines)

        indent_level = 0
        formatted_lines = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                formatted_lines.append('')
                continue

            if stripped.startswith('}'):
                indent_level = max(0, indent_level - 1)

            if stripped.startswith('#'):
                formatted_lines.append(stripped)
                continue

            if re.match(r'^(case\b.*:|default\s*:|public\s*:|private\s*:|protected\s*:)', stripped):
                label_indent = max(0, indent_level - 1)
                formatted_lines.append(('    ' * label_indent) + stripped)
                continue

            formatted_lines.append(('    ' * indent_level) + stripped)

            open_braces = stripped.count('{')
            close_braces = stripped.count('}')
            if stripped.startswith('}'):
                diff = open_braces - (close_braces - 1)
            else:
                diff = open_braces - close_braces

            indent_level = max(0, indent_level + diff)

        return '\n'.join(formatted_lines)

    @classmethod
    def auto_indent_sql(cls, code: str) -> str:
        lines = [l.strip() for l in code.split('\n') if l.strip()]
        if not lines:
            return code
        major_clauses = re.compile(r'^(SELECT|FROM|WHERE|GROUP\s+BY|ORDER\s+BY|HAVING|LIMIT|INSERT\s+INTO|VALUES|UPDATE|SET|DELETE\s+FROM|CREATE\s+TABLE|ALTER\s+TABLE)\b', re.IGNORECASE)
        formatted = []
        indent = 0
        for line in lines:
            if major_clauses.match(line):
                indent = 0
                formatted.append(line)
                if re.match(r'^(SELECT|SET|CREATE\s+TABLE)\b', line, re.IGNORECASE):
                    indent = 1
            elif re.match(r'^(AND|OR|JOIN|LEFT\s+JOIN|RIGHT\s+JOIN|INNER\s+JOIN|ON)\b', line, re.IGNORECASE):
                formatted.append('    ' + line)
            else:
                formatted.append(('    ' * indent) + line)
        return '\n'.join(formatted)

    @classmethod
    def auto_indent_html(cls, code: str) -> str:
        lines = [l.strip() for l in code.split('\n') if l.strip()]
        if not lines:
            return code
        void_tags = {'area', 'base', 'br', 'col', 'embed', 'hr', 'img', 'input', 'link', 'meta', 'param', 'source', 'track', 'wbr'}
        formatted = []
        indent = 0
        for line in lines:
            closing = re.findall(r'<\s*/\s*([a-zA-Z0-9]+)', line)
            opening = re.findall(r'<\s*([a-zA-Z0-9]+)(?:[^>]*)(?<!/)>', line)
            if line.startswith('</'):
                indent = max(0, indent - 1)
            formatted.append(('    ' * indent) + line)
            open_count = sum(1 for t in opening if t.lower() not in void_tags)
            close_count = len(closing)
            delta = open_count - (close_count - (1 if line.startswith('</') else 0))
            indent = max(0, indent + delta)
        return '\n'.join(formatted)

    @classmethod
    def auto_indent_css(cls, code: str) -> str:
        lines = [l.strip() for l in code.split('\n') if l.strip()]
        if not lines:
            return code
        formatted = []
        indent = 0
        for line in lines:
            if line.startswith('}'):
                indent = max(0, indent - 1)
            formatted.append(('    ' * indent) + line)
            open_count = line.count('{')
            close_count = line.count('}')
            if line.startswith('}'):
                indent = max(0, indent + open_count - (close_count - 1))
            else:
                indent = max(0, indent + open_count - close_count)
        return '\n'.join(formatted)

    @classmethod
    def format_code(cls, code: str, language: str = 'python') -> str:
        if not code or not code.strip():
            return code
        lang = (language or 'python').lower()
        code = cls.clean_and_dedent(code)
        code = cls.repair_first_line_stripped(code, lang)
        if lang in ['c++', 'cplusplus', 'c', 'cpp', 'java', 'js', 'javascript']:
            return cls.auto_indent_cpp(code)
        elif lang in ['python', 'py', 'python3']:
            return cls.auto_indent_python(code)
        elif lang in ['sql']:
            return cls.auto_indent_sql(code)
        elif lang in ['html', 'markup', 'html5']:
            return cls.auto_indent_html(code)
        elif lang in ['css', 'css3']:
            return cls.auto_indent_css(code)
        return code


class SmartCodeDetector:
    """
    AI-Heuristic Automatic Code Detection & Language Classifier
    Detects Python, C++, SQL, and HTML in raw exam text without explicit backticks or tags.
    """

    CPP_SIGNATURES = [
        r'#include\s*<', r'using\s+namespace\s+std', r'\bint\s+main\s*\(',
        r'\bcout\s*<<', r'\bcin\s*>>', r'\bstd::', r'\bvector\s*<', r'\bqueue\s*<',
        r'\bstack\s*<', r'\bpair\s*<', r'\bmap\s*<', r'\bset\s*<',
        r'\bprintf\s*\(', r'\bscanf\s*\(', r'\bvoid\s+\w+\s*\(',
        r'\bfor\s*\(\s*(?:int\s+)?', r'\bwhile\s*\(', r';\s*$', r'\{\s*$', r'^\s*\}\s*$',
        r'\bendl\b', r'\bnullptr\b', r'\bstruct\s+\w+\s*\{', r'\bclass\s+\w+\s*\{',
        r'\breturn\s+0\s*;'
    ]

    PYTHON_SIGNATURES = [
        r'\bdef\s+\w+\s*\(', r'\bprint\s*\(', r'\bimport\s+\w+', r'\bfrom\s+\w+\s+import',
        r'\bfor\s+\w+\s+in\s+range\s*\(', r'\bfor\s+\w+\s+in\s+', r'\bwhile\s+.*:',
        r'\bif\s+.+:\s*$', r'\belif\s+.+:\s*$', r'\belse\s*:\s*$', r'\breturn\b',
        r'\b\[\s*x\s+for\s+x\s+in\s+', r'\blen\s*\(', r'\brange\s*\(', r'\bsum\s*\(',
        r'\bappend\s*\(', r'\.split\s*\(', r'\binput\s*\(', r'\blambda\s+',
        r'\btry\s*:\s*$', r'\bexcept\b', r'\bclass\s+\w+\s*(?:\(.*?\))?:'
    ]

    SQL_SIGNATURES = [
        r'\bSELECT\s+', r'\bFROM\s+\w+', r'\bWHERE\s+', r'\bINNER\s+JOIN\b',
        r'\bLEFT\s+JOIN\b', r'\bRIGHT\s+JOIN\b', r'\bGROUP\s+BY\b', r'\bORDER\s+BY\b',
        r'\bINSERT\s+INTO\b', r'\bUPDATE\s+\w+\s+SET\b', r'\bDELETE\s+FROM\b',
        r'\bCREATE\s+TABLE\b', r'\bPRIMARY\s+KEY\b', r'\bFOREIGN\s+KEY\b',
        r'\bCOUNT\s*\(', r'\bSUM\s*\(', r'\bAVG\s*\(', r'\bMAX\s*\(', r'\bMIN\s*\(',
        r'\bHAVING\s+', r'\bDISTINCT\s+', r'\bVALUES\s*\('
    ]

    HTML_SIGNATURES = [
        r'<!DOCTYPE\s+html>', r'<\s*html\b', r'<\s*head\b', r'<\s*body\b',
        r'<\s*div\b', r'<\s*span\b', r'<\s*table\b', r'<\s*tr\b', r'<\s*td\b', r'<\s*th\b',
        r'<\s*h[1-6]\b', r'<\s*p\b', r'<\s*a\s+href=', r'<\s*img\s+src=', r'<\s*script\b',
        r'<\s*style\b', r'<\s*\/\s*\w+\s*>', r'<\s*video\b', r'<\s*audio\b', r'<\s*form\b',
        r'<\s*input\b', r'<\s*button\b', r'<\s*link\s+href='
    ]

    CSS_SIGNATURES = [
        r'(?:^[a-zA-Z0-9_\-\.#\s,>+~*:]+)\s*\{[\s\S]*?(?:color|margin|padding|background|font|display|width|height|border|text|position|top|left|right|bottom|float|clear|flex|align|justify|opacity|z-index)\s*:',
        r'^\s*(?:\.[a-zA-Z_-][a-zA-Z0-9_-]*|#[a-zA-Z_-][a-zA-Z0-9_-]*|body\b|html\b|p\b|h[1-6]\b|div\b|table\b|tr\b|td\b|th\b|ul\b|ol\b|li\b|span\b|a\b|img\b|button\b|input\b|form\b|header\b|footer\b|nav\b|section\b|article\b|main\b)\s*\{',
        r'^\s*[a-zA-Z0-9_\-\.#\s,>+~*:]+:[a-zA-Z_-]+\s*\{',
        r'^\s*@(media|keyframes|import|font-face|supports)\b',
        r'^\s*color\s*:\s*[^;{}]+;?',
        r'^\s*background(?:-color|-image|-size|-position|-repeat)?\s*:\s*[^;{}]+;?',
        r'^\s*font(?:-size|-family|-weight|-style)?\s*:\s*[^;{}]+;?',
        r'^\s*text-(?:align|decoration|transform|indent|shadow)\s*:\s*[^;{}]+;?',
        r'^\s*margin(?:-(?:top|bottom|left|right))?\s*:\s*[^;{}]+;?',
        r'^\s*padding(?:-(?:top|bottom|left|right))?\s*:\s*[^;{}]+;?',
        r'^\s*border(?:-(?:top|bottom|left|right|radius|color|width|style))?\s*:\s*[^;{}]+;?',
        r'^\s*display\s*:\s*(?:flex|grid|block|inline|inline-block|none|table)\s*;?',
        r'^\s*(?:width|height|max-width|min-width|max-height|min-height)\s*:\s*[^;{}]+;?',
        r'^\s*position\s*:\s*(?:relative|absolute|fixed|sticky|static)\s*;?',
        r'^\s*(?:top|bottom|left|right|z-index)\s*:\s*[^;{}]+;?',
        r'^\s*(?:justify-content|align-items|align-content|flex-direction|flex-wrap)\s*:\s*[^;{}]+;?',
        r'^\s*(?:float|clear|overflow|box-sizing|opacity|cursor|line-height|letter-spacing)\s*:\s*[^;{}]+;?',
        r'^\s*[a-z][a-z0-9\-]*\s*:\s*[^;{}]+\s*;\s*$'
    ]

    LANG_HEADERS = {
        'python': re.compile(r'^\s*(?:\*\*)?(?:\[\s*)?(?:(?:Chương\s*trình|Đoạn\s*(?:mã|code|chương\s*trình)|Mã\s*nguồn)(?:\s+viết)?(?:\s+bằng)?\s*)?(?:ngôn\s*ngữ\s*)?(?:Python|Py|Python\s*3)(?:\s*\])?(?:\*\*)?\s*[:\.\-]?\s*$', re.IGNORECASE),
        'cpp': re.compile(r'^\s*(?:\*\*)?(?:\[\s*)?(?:(?:Chương\s*trình|Đoạn\s*(?:mã|code|chương\s*trình)|Mã\s*nguồn)(?:\s+viết)?(?:\s+bằng)?\s*)?(?:ngôn\s*ngữ\s*)?(?:C\+\+|CPP|C|CPlusPlus)(?:\s*\])?(?:\*\*)?\s*[:\.\-]?\s*$', re.IGNORECASE),
        'sql': re.compile(r'^\s*(?:\*\*)?(?:\[\s*)?(?:(?:Chương\s*trình|Đoạn\s*(?:mã|code)|Mã\s*nguồn)(?:\s+viết)?(?:\s+bằng)?\s*)?(?:ngôn\s*ngữ\s*)?(?:SQL|RDBMS|Truy\s*vấn\s*SQL)(?:\s*\])?(?:\*\*)?\s*[:\.\-]?\s*$', re.IGNORECASE),
        'html': re.compile(r'^\s*(?:\*\*)?(?:\[\s*)?(?:(?:Chương\s*trình|Đoạn\s*(?:mã|code)|Mã\s*nguồn)(?:\s+viết)?(?:\s+bằng)?\s*)?(?:ngôn\s*ngữ\s*)?(?:HTML|HTML5|Mã\s*HTML)(?:\s*\])?(?:\*\*)?\s*[:\.\-]?\s*$', re.IGNORECASE),
        'css': re.compile(r'^\s*(?:\*\*)?(?:\[\s*)?(?:(?:Chương\s*trình|Đoạn\s*(?:mã|code)|Mã\s*nguồn)(?:\s+viết)?(?:\s+bằng)?\s*)?(?:ngôn\s*ngữ\s*)?(?:CSS|CSS3|Mã\s*CSS|Định\s*dạng\s*CSS)(?:\s*\])?(?:\*\*)?\s*[:\.\-]?\s*$', re.IGNORECASE),
    }

    CPP_START_KEYWORDS = re.compile(r'^\s*(#include\s*<|using\s+namespace|int\s+main|void\s+|int\s+\w+\s*\(|template\s*<|struct\s+|class\s+)')
    PYTHON_START_KEYWORDS = re.compile(r'^\s*(def\s+|import\s+|from\s+\w+\s+import|\w+\s*=\s*\[|for\s+\w+\s+in\s+|print\s*\(|class\s+\w+)')
    SQL_START_KEYWORDS = re.compile(r'^\s*(SELECT\s+|INSERT\s+INTO|UPDATE\s+|DELETE\s+FROM|CREATE\s+TABLE)', re.IGNORECASE)
    HTML_START_KEYWORDS = re.compile(r'^\s*(<!DOCTYPE\s+html>|<\s*html\b|<\s*div\b|<\s*table\b|<\s*p\b|<\s*style\b)', re.IGNORECASE)
    CSS_START_KEYWORDS = re.compile(r'^\s*(@(?:media|keyframes|import|font-face|supports)|(?:[a-zA-Z0-9_\-\.#\s,>+~*:]+)\s*\{\s*(?:color|margin|padding|background|font|display|width|height|border|text|position|top|left|float)\s*:|(?:\.[a-zA-Z_-]|#[a-zA-Z_-]|body\b|html\b|p\b|h[1-6]\b|div\b|table\b|span\b|a\b|ul\b|li\b)\s*\{)', re.IGNORECASE)

    @classmethod
    def check_header(cls, line: str) -> str:
        for lang, pat in cls.LANG_HEADERS.items():
            if pat.match(line.strip()):
                return lang
        return ''

    @classmethod
    def split_code_lines(cls, lines: List[str]) -> List[Dict[str, str]]:
        snippets: List[Dict[str, str]] = []
        current_lang: str = ''
        current_block: List[str] = []

        for line in lines:
            line_str = line.strip()
            if not line_str:
                if current_block:
                    current_block.append(line)
                continue

            header_lang = cls.check_header(line_str)
            if header_lang:
                if current_block and current_lang:
                    code_text = "\n".join(current_block).rstrip()
                    if code_text.strip():
                        code_text = SmartCodeFormatter.format_code(code_text, current_lang)
                        snippets.append({'language': current_lang, 'code': code_text})
                current_lang = header_lang
                current_block = []
                continue

            if current_lang == 'python' and cls.CPP_START_KEYWORDS.match(line):
                code_text = "\n".join(current_block).rstrip()
                if code_text.strip():
                    code_text = SmartCodeFormatter.format_code(code_text, 'python')
                    snippets.append({'language': 'python', 'code': code_text})
                current_lang = 'cpp'
                current_block = [line]
                continue

            if current_lang == 'cpp' and cls.PYTHON_START_KEYWORDS.match(line) and not line_str.endswith(';'):
                code_text = "\n".join(current_block).rstrip()
                if code_text.strip():
                    code_text = SmartCodeFormatter.format_code(code_text, 'cpp')
                    snippets.append({'language': 'cpp', 'code': code_text})
                current_lang = 'python'
                current_block = [line]
                continue

            if current_lang == 'html' and cls.CSS_START_KEYWORDS.match(line):
                code_text = "\n".join(current_block).rstrip()
                if code_text.strip():
                    code_text = SmartCodeFormatter.format_code(code_text, 'html')
                    snippets.append({'language': 'html', 'code': code_text})
                current_lang = 'css'
                current_block = [line]
                continue

            if current_lang == 'css' and cls.HTML_START_KEYWORDS.match(line):
                code_text = "\n".join(current_block).rstrip()
                if code_text.strip():
                    code_text = SmartCodeFormatter.format_code(code_text, 'css')
                    snippets.append({'language': 'css', 'code': code_text})
                current_lang = 'html'
                current_block = [line]
                continue

            if not current_lang:
                if cls.CPP_START_KEYWORDS.match(line):
                    current_lang = 'cpp'
                elif cls.PYTHON_START_KEYWORDS.match(line):
                    current_lang = 'python'
                elif cls.SQL_START_KEYWORDS.match(line):
                    current_lang = 'sql'
                elif cls.HTML_START_KEYWORDS.match(line):
                    current_lang = 'html'
                elif cls.CSS_START_KEYWORDS.match(line):
                    current_lang = 'css'

            current_block.append(line)

        if current_block:
            code_text = "\n".join(current_block).rstrip()
            if code_text.strip():
                if not current_lang:
                    current_lang = cls.classify_language(code_text)
                code_text = SmartCodeFormatter.format_code(code_text, current_lang)
                snippets.append({'language': current_lang, 'code': code_text})

        return snippets

    @classmethod
    def classify_language(cls, code_str: str) -> str:
        """
        Determines programming language based on syntax score heuristics.
        Returns: 'python', 'cpp', 'sql', 'html', or 'css'.
        """
        if not code_str or not code_str.strip():
            return 'python'

        scores = {'python': 0, 'cpp': 0, 'sql': 0, 'html': 0, 'css': 0}

        for pattern in cls.PYTHON_SIGNATURES:
            if re.search(pattern, code_str, re.IGNORECASE if 'for' not in pattern else 0):
                scores['python'] += 2

        for pattern in cls.CPP_SIGNATURES:
            if re.search(pattern, code_str, re.IGNORECASE if 'for' not in pattern else 0):
                scores['cpp'] += 2

        for pattern in cls.SQL_SIGNATURES:
            if re.search(pattern, code_str, re.IGNORECASE):
                scores['sql'] += 3

        for pattern in cls.HTML_SIGNATURES:
            if re.search(pattern, code_str, re.IGNORECASE):
                scores['html'] += 3

        for pattern in cls.CSS_SIGNATURES:
            if re.search(pattern, code_str, re.IGNORECASE | re.MULTILINE):
                scores['css'] += 2

        # Count CSS declarations (property: value;)
        css_props = re.findall(r'(?:\b(?:color|background|font|margin|padding|border|display|width|height|text|align|justify|flex|position|top|left|right|bottom|opacity|overflow|z-index)[a-z\-]*\s*:\s*[^;{}]+;)', code_str, re.IGNORECASE)
        scores['css'] += len(css_props) * 2

        # Check for CSS selector blocks like selector { ... }
        if re.search(r'(?:^[a-zA-Z0-9_\-\.#\s,>+~*:]+)\s*\{[^}]*\}', code_str, re.MULTILINE):
            scores['css'] += 3

        semicolon_count = code_str.count(';')
        lines_count = max(1, len(code_str.split('\n')))
        if semicolon_count / lines_count > 0.4:
            scores['cpp'] += 2
            scores['sql'] += 1
            if scores['css'] > 0:
                scores['css'] += 2

        if re.search(r'\n\s{4,}\w+', code_str) and not re.search(r'[\{\}]', code_str):
            scores['python'] += 2

        # If it has strong CSS declarations and no C++ constructs like #include, cout, int main, types with assignments
        if scores['css'] > 0 and not re.search(r'(#include|using\s+namespace|int\s+main|std::|cout|cin|\bint\s+\w+\s*=)', code_str):
            if not re.search(r'<!DOCTYPE|<\s*html|<\s*body|<\s*head|<\s*table|<\s*div\b|<\s*p\b|<\s*script\b', code_str, re.IGNORECASE):
                scores['css'] += 3

        best_lang = max(scores, key=scores.get)
        return best_lang if scores[best_lang] > 0 else 'python'

    @classmethod
    def is_line_code(cls, line: str) -> bool:
        line_s = line.strip()
        if not line_s:
            return False

        if cls.check_header(line_s):
            return True

        if re.search(r'^(Câu|PHẦN|A\.|B\.|C\.|D\.|a\)|b\)|c\)|d\)|\[)', line_s, re.IGNORECASE):
            return False

        all_patterns = cls.CPP_SIGNATURES + cls.PYTHON_SIGNATURES + cls.SQL_SIGNATURES + cls.HTML_SIGNATURES + cls.CSS_SIGNATURES
        for p in all_patterns:
            if re.search(p, line_s, re.IGNORECASE):
                return True

        if re.search(r'^\s*[a-zA-Z_]\w*\s*(=|\+=|-=|\*=|\/=|%=)\s*[^?]+$', line_s):
            return True

        if re.search(r'^\s*[a-zA-Z_]\w*\s*\([^\)]*\)\s*;?\s*$', line_s):
            return True

        return False

    @classmethod
    def extract_unformatted_code(cls, lines: List[str]) -> Tuple[List[str], str, str]:
        """
        Extracts untagged code lines and automatically splits multi-language blocks (Python + C++).
        Returns: (filtered_lines, code_snippet, code_language)
        """
        code_start = -1
        code_end = -1

        for idx, l in enumerate(lines):
            if cls.is_line_code(l):
                if code_start == -1:
                    code_start = idx
                code_end = idx

        if code_start != -1 and code_end != -1 and (code_end - code_start + 1) >= 1:
            code_lines = lines[code_start:code_end + 1]
            segments = cls.split_code_lines(code_lines)

            if len(segments) > 1:
                combined_code = "\n\n".join([f"```{s['language']}\n{s['code']}\n```" for s in segments])
                primary_lang = segments[0]['language']
            elif len(segments) == 1:
                combined_code = segments[0]['code']
                primary_lang = segments[0]['language']
            else:
                combined_code = SmartCodeFormatter.clean_and_dedent("\n".join(code_lines))
                primary_lang = cls.classify_language(combined_code)
                combined_code = SmartCodeFormatter.format_code(combined_code, primary_lang)

            filtered = lines[:code_start] + lines[code_end + 1:]
            return filtered, combined_code, primary_lang

        return lines, "", "python"


class DocxExamParser:
    @staticmethod
    def is_run_red(run) -> bool:
        """
        Detects if a docx Run has red font color or red highlight.
        """
        try:
            # 1. Check font.color.rgb
            if run.font and run.font.color and run.font.color.rgb:
                rgb = run.font.color.rgb
                if rgb[0] >= 180 and rgb[1] < 120 and rgb[2] < 120:
                    return True

            # 2. Check XML w:color or highlight
            if hasattr(run, '_r'):
                for child in run._r.iter():
                    tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                    if tag == 'color':
                        val = child.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') or child.get('val') or ''
                        val = val.lower()
                        if val == 'red' or any(val.startswith(prefix) for prefix in ['ff', 'c0', 'ed', 'e5', 'dc', 'e0', 'ef', 'f8', 'f4', 'eb', 'd3']):
                            return True
                    if tag == 'highlight':
                        val = child.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') or child.get('val') or ''
                        if val.lower() == 'red':
                            return True
        except Exception:
            pass
        return False

    @staticmethod
    def extract_images_from_docx(doc: docx.Document) -> Dict[str, str]:
        """
        Extracts embedded images from docx document and saves to media/exam_images/
        Returns a mapping from rId to public media URL.
        """
        image_mapping = {}
        try:
            media_root = getattr(settings, 'MEDIA_ROOT', os.path.join(os.path.dirname(os.path.dirname(__file__)), 'media'))
            media_url = getattr(settings, 'MEDIA_URL', '/media/')
        except Exception:
            media_root = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'media')
            media_url = '/media/'

        images_dir = os.path.join(media_root, 'exam_images')
        os.makedirs(images_dir, exist_ok=True)

        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.target_ref:
                try:
                    image_part = rel.target_part
                    ext = image_part.content_type.split('/')[-1]
                    if ext == 'jpeg': ext = 'jpg'
                    filename = f"img_{uuid.uuid4().hex[:12]}.{ext}"
                    filepath = os.path.join(images_dir, filename)
                    with open(filepath, "wb") as f:
                        f.write(image_part.blob)
                    image_mapping[rel_id] = f"{media_url}exam_images/{filename}"
                except Exception as e:
                    print(f"Warning extracting image {rel_id}: {e}")

        return image_mapping

    @classmethod
    def _process_paragraph(cls, p, image_mapping: Dict[str, str]) -> str:
        def _get_image_rids(element):
            rids = []
            if element is None:
                return rids
            for child in element.iter():
                tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                if tag == 'blip':
                    rid = child.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if rid:
                        rids.append(rid)
                elif tag == 'imagedata':
                    rid = child.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    if rid:
                        rids.append(rid)
            return rids

        p_text = ""
        has_red_in_p = False

        # Extract Word paragraph indentation (left_indent or first_line_indent)
        indent_spaces = ""
        try:
            pf = p.paragraph_format
            total_indent_pt = 0.0
            if pf.left_indent and pf.left_indent.pt > 0:
                total_indent_pt += pf.left_indent.pt
            if pf.first_line_indent and pf.first_line_indent.pt > 0:
                total_indent_pt += pf.first_line_indent.pt
            if total_indent_pt >= 12:  # at least ~1/6 inch
                level = max(1, int(round(total_indent_pt / 36.0)))
                indent_spaces = "    " * level
        except Exception:
            pass

        for run in p.runs:
            r_text = run.text

            # Check for run-level images
            if hasattr(run, '_r'):
                for rid in _get_image_rids(run._r):
                    if rid in image_mapping:
                        p_text += f"\n![Hình ảnh]({image_mapping[rid]})\n"

            if not r_text: continue
            # Convert tabs to 4 spaces, non-breaking space \xa0 to regular space
            r_text = r_text.replace('\t', '    ').replace('\xa0', ' ')

            if cls.is_run_red(run):
                has_red_in_p = True
                p_text += f"[RED]{r_text}[/RED]"
            else:
                p_text += r_text

        # Check if paragraph itself has drawing images outside runs
        if hasattr(p, '_p'):
            for rid in _get_image_rids(p._p):
                if rid in image_mapping and f"![Hình ảnh]({image_mapping[rid]})" not in p_text:
                    p_text += f"\n![Hình ảnh]({image_mapping[rid]})\n"

        if indent_spaces:
            if '\n' in p_text:
                p_text = '\n'.join((indent_spaces + line if line.strip() else line) for line in p_text.split('\n'))
            else:
                p_text = indent_spaces + p_text
        if p_text and has_red_in_p:
            p_text = re.sub(r'\[RED\]\s*\*?([A-Da-d]\*?)\s*[\.:\)]\s*\[/RED\]', r'*\1. ', p_text)
            p_text = re.sub(r'^\s*([A-Da-d]\*?)\s*[\.:\)]\s*\[RED\]', r'*\1. ', p_text)
            p_text = p_text.replace('[RED]', '').replace('[/RED]', '')

        # Preserve leading whitespace! Only strip trailing newlines and spaces.
        return p_text.rstrip('\r\n ')

    @classmethod
    def _is_answer_table(cls, tbl: docx.table.Table) -> bool:
        rows = tbl.rows
        if len(rows) < 2:
            return False
        all_cell_texts = [[c.text.strip() for c in r.cells] for r in rows]
        for r_idx in range(len(rows) - 1):
            r1 = all_cell_texts[r_idx]
            r2 = all_cell_texts[r_idx + 1]
            cau_count = sum(1 for c in r1 if re.search(r'^(?:Câu\s*\d+|\d+)$', c, re.IGNORECASE))
            ans_count = sum(1 for c in r2 if re.search(r'^[A-D]$', c, re.IGNORECASE))
            if cau_count >= 2 and ans_count >= 2:
                return True
            if any('đáp án' in c.lower() for c in r1) or any('bảng đáp án' in c.lower() for c in r1):
                return True
        return False

    @classmethod
    def _process_content_table(cls, tbl: docx.table.Table, image_mapping: Dict[str, str]) -> List[str]:
        rows = tbl.rows
        if not rows:
            return []

        num_cols = len(tbl.columns)
        num_rows = len(rows)

        # Build 2D grid of cell contents
        grid: List[List[str]] = []
        for r in rows:
            row_texts = []
            for c in r.cells:
                cell_parts = []
                for p in c.paragraphs:
                    p_t = cls._process_paragraph(p, image_mapping)
                    if p_t:
                        cell_parts.append(p_t)
                row_texts.append("\n".join(cell_parts).strip())
            grid.append(row_texts)

        # 1. Check if table is a Code Comparison Table (e.g. Column 0 = Python, Column 1 = C++)
        col_snippets = []
        is_code_table = False

        for col_idx in range(num_cols):
            col_cells = [grid[r_idx][col_idx] for r_idx in range(num_rows) if col_idx < len(grid[r_idx])]
            full_col_text = "\n".join([c for c in col_cells if c.strip()])
            if not full_col_text.strip():
                continue

            code_lines = []
            col_lang = ''

            for cell_text in col_cells:
                if not cell_text.strip():
                    continue
                for line in cell_text.split('\n'):
                    l_str = line.strip()
                    if not l_str:
                        if code_lines:
                            code_lines.append(line)
                        continue
                    line_h = SmartCodeDetector.check_header(l_str)
                    if line_h:
                        if not col_lang:
                            col_lang = line_h
                        continue
                    code_lines.append(line)

            col_code = "\n".join(code_lines).strip()
            if col_code:
                detected_lang = col_lang or SmartCodeDetector.classify_language(col_code)
                has_code_syntax = any(SmartCodeDetector.is_line_code(line) for line in col_code.split('\n'))
                if col_lang or has_code_syntax:
                    is_code_table = True
                    formatted_code = SmartCodeFormatter.format_code(col_code, detected_lang)
                    col_snippets.append({
                        'language': detected_lang,
                        'code': formatted_code
                    })

        if is_code_table and col_snippets:
            res_lines = []
            for snip in col_snippets:
                res_lines.append(f"```{snip['language']}\n{snip['code']}\n```")
            return res_lines

        # 2. Otherwise format as Markdown Table
        md_lines = []
        if num_rows >= 1:
            header_cells = grid[0]
            md_lines.append("| " + " | ".join(header_cells) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")
            for r_idx in range(1, num_rows):
                md_lines.append("| " + " | ".join(grid[r_idx]) + " |")
        return md_lines

    @classmethod
    def parse_docx_file(cls, file_path_or_buffer) -> Dict[str, Any]:
        """
        Parses a .docx exam file with Image Extraction, Run-level Red font detection,
        Table traversal (Code tables & Data tables), Table Answer Keys, and Part II Common + CS/ICT.
        """
        doc = docx.Document(file_path_or_buffer)
        image_mapping = cls.extract_images_from_docx(doc)
        raw_lines: List[str] = []

        # Traverse document elements in sequential document order (paragraphs and tables)
        for child in doc.element.body:
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag == 'p':
                p = docx.text.paragraph.Paragraph(child, doc)
                p_text = cls._process_paragraph(p, image_mapping)
                if p_text:
                    raw_lines.append(p_text)
            elif tag == 'tbl':
                tbl = docx.table.Table(child, doc)
                if cls._is_answer_table(tbl):
                    # Skip answer key table from question body text
                    continue
                else:
                    table_lines = cls._process_content_table(tbl, image_mapping)
                    for t_line in table_lines:
                        if t_line.strip():
                            raw_lines.append(t_line)

        # Also extract table answer keys if any
        table_keys = cls._extract_answer_tables(doc.tables)

        full_text = "\n".join(raw_lines)
        parsed = cls.parse_raw_text(full_text)

        # Apply Table Answer Keys backfill if any answers were missing
        cls._apply_table_answer_keys(parsed, table_keys)

        # Provide raw_text for frontend text area sync
        parsed['raw_text'] = full_text

        return parsed

    @classmethod
    def _extract_pdf_page_text(cls, page) -> str:
        """
        Extracts text from a pdfplumber page while preserving code line indentation.
        """
        try:
            words = page.extract_words(y_tolerance=3, x_tolerance=3)
            if not words:
                return page.extract_text() or ''

            lines = []
            current_line = []
            current_top = None

            sorted_words = sorted(words, key=lambda w: (round(w['top'] / 3) * 3, w['x0']))

            for w in sorted_words:
                if current_top is None or abs(w['top'] - current_top) <= 3:
                    current_line.append(w)
                    if current_top is None:
                        current_top = w['top']
                else:
                    lines.append(current_line)
                    current_line = [w]
                    current_top = w['top']
            if current_line:
                lines.append(current_line)

            x0_values = [line[0]['x0'] for line in lines if line]
            if not x0_values:
                return page.extract_text() or ''

            sorted_x0 = sorted(x0_values)
            base_x0 = sorted_x0[max(0, int(len(sorted_x0) * 0.05))]

            result_lines = []
            for line in lines:
                line_x0 = line[0]['x0']
                diff = line_x0 - base_x0
                indent_spaces = ''
                if diff >= 12:
                    level = max(1, int(round((diff / 7.2) / 4.0)))
                    indent_spaces = '    ' * level
                line_text = indent_spaces + ' '.join(w['text'] for w in line)
                result_lines.append(line_text)

            return '\n'.join(result_lines)
        except Exception:
            return page.extract_text() or ''

    @classmethod
    def parse_pdf_file(cls, file_path_or_buffer) -> Dict[str, Any]:
        """
        Parses an exam PDF file using pdfplumber, extracts text and embedded images,
        and converts it into structured exam questions.
        """
        import pdfplumber

        try:
            media_root = getattr(settings, 'MEDIA_ROOT', os.path.join(os.path.dirname(os.path.dirname(__file__)), 'media'))
            media_url = getattr(settings, 'MEDIA_URL', '/media/')
        except Exception:
            media_root = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'media')
            media_url = '/media/'

        images_dir = os.path.join(media_root, 'exam_images')
        os.makedirs(images_dir, exist_ok=True)

        page_texts = []
        with pdfplumber.open(file_path_or_buffer) as pdf:
            for page_idx, page in enumerate(pdf.pages):
                p_text = cls._extract_pdf_page_text(page)
                # Extract embedded images on this page if any
                try:
                    for img_idx, img in enumerate(page.images):
                        stream = img.get('stream')
                        if stream and hasattr(stream, 'get_data'):
                            raw_data = stream.get_data()
                            ext = 'png'
                            filename = f"pdf_img_{page_idx+1}_{img_idx+1}_{uuid.uuid4().hex[:8]}.{ext}"
                            filepath = os.path.join(images_dir, filename)
                            with open(filepath, 'wb') as f:
                                f.write(raw_data)
                            img_url = f"{media_url}exam_images/{filename}"
                            p_text += f"\n![Hình ảnh]({img_url})\n"
                except Exception as img_err:
                    print(f"Warning extracting PDF image on page {page_idx}: {img_err}")

                page_texts.append(p_text)

        full_text = "\n".join(page_texts)
        parsed = cls.parse_raw_text(full_text)
        parsed['raw_text'] = full_text
        return parsed

    @classmethod
    def _extract_answer_tables(cls, tables) -> Dict[str, Any]:
        """
        Extracts answers from Answer Key Tables at the end of the document.
        """
        p1_keys: Dict[int, str] = {}
        p2_branch_keys: Dict[int, Dict[str, bool]] = {}

        for table in tables:
            rows = table.rows
            if len(rows) < 2: continue

            all_cell_texts = [[c.text.strip() for c in r.cells] for r in rows]

            # Check if this is a Part I table (Row 0: Câu 1 2 3... Row 1: Chọn B D C...)
            for r_idx in range(0, len(rows) - 1, 2):
                header_row = all_cell_texts[r_idx]
                val_row = all_cell_texts[r_idx + 1]

                if any('Câu' in c for c in header_row) and any(re.search(r'^[A-D]$', c, re.IGNORECASE) for c in val_row):
                    for c_idx in range(len(header_row)):
                        q_str = header_row[c_idx]
                        v_str = val_row[c_idx] if c_idx < len(val_row) else ''
                        q_match = re.search(r'(\d+)', q_str)
                        v_match = re.search(r'([A-D])', v_str, re.IGNORECASE)
                        if q_match and v_match:
                            q_num = int(q_match.group(1))
                            p1_keys[q_num] = v_match.group(1).upper()

            # Check Part II table
            for row in rows:
                for c_idx, cell in enumerate(row.cells):
                    c_text = cell.text.strip()
                    if re.search(r'[a-d]\)\s*(Đúng|Sai|Dung|Đ|S)', c_text, re.IGNORECASE):
                        sub_dict = {}
                        for line in c_text.split('\n'):
                            m = re.search(r'([a-d])\)\s*(Đúng|Sai|Dung|Đ|S)', line, re.IGNORECASE)
                            if m:
                                label = m.group(1).lower()
                                is_true = bool(re.search(r'Đúng|Dung|Đ', m.group(2), re.IGNORECASE))
                                sub_dict[label] = is_true
                        if sub_dict:
                            p2_branch_keys[c_idx] = sub_dict

        return {
            'part1_keys': p1_keys,
            'p2_branch_keys': p2_branch_keys,
        }

    @classmethod
    def _apply_table_answer_keys(cls, parsed: Dict[str, Any], table_keys: Dict[str, Any]):
        """
        Backfills question options from table keys if not already marked.
        """
        p1_keys = table_keys.get('part1_keys', {})
        if not p1_keys: return

        for q in parsed.get('questions', []):
            if q['part_type'] == 'PART_I':
                q_num = q['order_index']
                if q_num in p1_keys:
                    correct_label = p1_keys[q_num]
                    has_marked = any(opt['is_correct'] for opt in q['options'])
                    if not has_marked:
                        for opt in q['options']:
                            if opt['label'].upper() == correct_label:
                                opt['is_correct'] = True

    @classmethod
    def parse_raw_text(cls, text: str) -> Dict[str, Any]:
        """
        Parses formatted exam text using Azota-compatible regex and state-machine tokenizer.
        """
        warnings: List[str] = []
        exam_metadata = {
            'title': 'ĐỀ THI TIN HỌC HSG THPT QUẤT LÂM',
            'duration_minutes': 50,
            'description': ''
        }

        # 1. Separate Main Exam content from Footer (e.g. ---HẾT--- and Bảng đáp án)
        exam_body = text
        footer_text = ""

        het_match = re.search(r'(-{3,}\s*(HẾT|HET)\s*-{3,}|Bảng\s+đáp\s+án|BẢNG\s+ĐÁP\s+ÁN)', text, re.IGNORECASE)
        if het_match:
            split_idx = het_match.start()
            exam_body = text[:split_idx]
            footer_text = text[split_idx:]

        # Extract answer keys from Footer (Azota format: 1A 2B 3C or Câu 4: a)Đ b)S c)S d)Đ)
        footer_p1_keys, footer_p2_keys = cls._parse_azota_footer_keys(footer_text)

        # 2. Parse Metadata Tags
        title_match = re.search(r'\[DE_THI\]\s*(.+)', exam_body, re.IGNORECASE)
        if title_match:
            exam_metadata['title'] = title_match.group(1).strip()

        duration_match = re.search(r'\[THOI_GIAN\]\s*(\d+)', exam_body, re.IGNORECASE)
        if duration_match:
            exam_metadata['duration_minutes'] = int(duration_match.group(1).strip())

        matrix_match = re.search(r'\[MA_TRAN\]\s*([A-Za-z0-9_]+)', exam_body, re.IGNORECASE)
        if matrix_match:
            exam_metadata['matrix_preset'] = matrix_match.group(1).strip()

        p1_total_match = re.search(r'\[(?:TONG_DIEM_P1|DIEM_P1|DIEM_PHAN_1)\]\s*([\d\.]+)', exam_body, re.IGNORECASE)
        if p1_total_match:
            exam_metadata['part1_total_points'] = float(p1_total_match.group(1).strip())

        p2_total_match = re.search(r'\[(?:TONG_DIEM_P2|DIEM_P2|DIEM_PHAN_2)\]\s*([\d\.]+)', exam_body, re.IGNORECASE)
        if p2_total_match:
            exam_metadata['part2_total_points'] = float(p2_total_match.group(1).strip())

        total_match = re.search(r'\[(?:TONG_DIEM|TONG_DIEM_DE)\]\s*([\d\.]+)', exam_body, re.IGNORECASE)
        if total_match:
            exam_metadata['total_points'] = float(total_match.group(1).strip())
        elif p1_total_match or p2_total_match:
            # L1: Dùng .get() tránh KeyError khi chỉ có Part 1 hoặc Part 2
            exam_metadata['total_points'] = round(exam_metadata.get('part1_total_points', 12.0) + exam_metadata.get('part2_total_points', 8.0), 2)

        # 3. Split into Sections
        lines = exam_body.split('\n')
        current_section = 'PART_I'
        current_branch = 'COMMON'

        section_lines = {
            ('PART_I', 'COMMON'): [],
            ('PART_II', 'COMMON'): [],
            ('PART_II', 'CS'): [],
            ('PART_II', 'ICT'): []
        }

        for line in lines:
            line_str = line.strip()
            if not line_str: continue

            # Section Header Detection (Azota format: "Phần 1. TRẮC NGHIỆM", "PHẦN II. Câu trắc nghiệm đúng sai...", "A. Phần chung...", "B. Phần riêng CS...", "C. Phần riêng ICT...")
            if not re.search(r'^\s*Câu\s*\d+', line_str, re.IGNORECASE):
                if re.search(r'\[PHAN_I\]|^\s*(?:PHẦN|PHAN)\s*(?:1|I)\b', line_str, re.IGNORECASE) and not re.search(r'PHẦN\s*II', line_str, re.IGNORECASE):
                    current_section = 'PART_I'
                    current_branch = 'COMMON'
                    continue
                elif re.search(r'\[PHAN_II_CS\]|\[PHAN_II_CHUYEN_DE_1\]|\[CS\]|^\s*(?:[A-C1-3]\.\s*)?(?:Chuyên\s*đề\s*Khoa\s*học\s*máy\s*tính|Khoa\s*học\s*máy\s*tính|.*(?:Khoa\s*học\s*máy\s*tính|định\s*hướng\s*CS|\bCS\b))', line_str, re.IGNORECASE):
                    current_section = 'PART_II'
                    current_branch = 'CS'
                    continue
                elif re.search(r'\[PHAN_II_ICT\]|\[PHAN_II_CHUYEN_DE_2\]|\[ICT\]|^\s*(?:[A-C1-3]\.\s*)?(?:Chuyên\s*đề\s*Tin\s*học\s*ứng\s*dụng|Tin\s*học\s*ứng\s*dụng|.*(?:Tin\s*học\s*ứng\s*dụng|định\s*hướng\s*ICT|\bICT\b))', line_str, re.IGNORECASE):
                    current_section = 'PART_II'
                    current_branch = 'ICT'
                    continue
                elif re.search(r'\[PHAN_II_CHUNG\]|^\s*(?:MỤC\s*A\b|[A-C1-3]\.\s*Phần\s*chung|Phần\s*chung\s*cho\s*tất\s*cả)', line_str, re.IGNORECASE):
                    current_section = 'PART_II'
                    current_branch = 'COMMON'
                    continue
                elif re.search(r'\[PHAN_II\]|^\s*(?:PHẦN|PHAN)\s*(?:2|II)\b', line_str, re.IGNORECASE):
                    current_section = 'PART_II'
                    current_branch = 'COMMON'
                    continue

            section_lines[(current_section, current_branch)].append(line)

        # 4. Parse Questions
        parsed_questions: List[Dict[str, Any]] = []

        # Part I Questions (Any number: 24, 30...)
        p1_questions = cls._parse_section_questions(
            lines=section_lines[('PART_I', 'COMMON')],
            part_type='PART_I',
            branch='COMMON',
            start_order_index=1,
            warnings=warnings
        )
        parsed_questions.extend(p1_questions)

        # Part II Common
        p2_common_questions = cls._parse_section_questions(
            lines=section_lines[('PART_II', 'COMMON')],
            part_type='PART_II',
            branch='COMMON',
            start_order_index=1,
            warnings=warnings
        )
        parsed_questions.extend(p2_common_questions)

        # Part II CS
        p2_cs_start_idx = len(p2_common_questions) + 1
        p2_cs_questions = cls._parse_section_questions(
            lines=section_lines[('PART_II', 'CS')],
            part_type='PART_II',
            branch='CS',
            start_order_index=p2_cs_start_idx,
            warnings=warnings
        )
        parsed_questions.extend(p2_cs_questions)

        # Part II ICT
        p2_ict_questions = cls._parse_section_questions(
            lines=section_lines[('PART_II', 'ICT')],
            part_type='PART_II',
            branch='ICT',
            start_order_index=p2_cs_start_idx,
            warnings=warnings
        )
        parsed_questions.extend(p2_ict_questions)

        # Apply Footer Keys if any question is missing key
        for q in parsed_questions:
            q_num = q['order_index']
            if q['part_type'] == 'PART_I' and q_num in footer_p1_keys:
                if not any(opt['is_correct'] for opt in q['options']):
                    for opt in q['options']:
                        if opt['label'].upper() == footer_p1_keys[q_num]:
                            opt['is_correct'] = True
            elif q['part_type'] == 'PART_II' and q_num in footer_p2_keys:
                # Sub-item boolean map
                sub_map = footer_p2_keys[q_num]
                for opt in q['options']:
                    opt_lbl = opt['label'].lower()
                    if opt_lbl in sub_map:
                        opt['is_correct'] = sub_map[opt_lbl]

        # Auto-detect matrix_preset and points if not explicitly defined
        p1_count = len(p1_questions)
        p2_eff = len(p2_common_questions) + max(len(p2_cs_questions), len(p2_ict_questions))
        p2_total_count = len(p2_common_questions) + len(p2_cs_questions) + len(p2_ict_questions)

        preset = exam_metadata.get('matrix_preset')
        if not preset or preset in ['HSG_QUAT_LAM', 'CUSTOM', '']:
            if p1_count >= 28 or p2_eff >= 5 or p2_total_count >= 7:
                preset = 'HSG_NINHBINH'
            elif (p1_count > 0 and p1_count <= 26) or p2_eff == 4 or p2_total_count == 6:
                preset = 'BGD_2025'
            else:
                preset = 'HSG_NINHBINH'
            exam_metadata['matrix_preset'] = preset

        if preset == 'HSG_NINHBINH':
            if 'part1_total_points' not in exam_metadata:
                exam_metadata['part1_total_points'] = 12.0
            if 'part2_total_points' not in exam_metadata:
                exam_metadata['part2_total_points'] = 8.0
            if 'total_points' not in exam_metadata:
                exam_metadata['total_points'] = 20.0
            # Apply default points if question didn't have explicit point tag
            for q in parsed_questions:
                if not q.get('has_explicit_point'):
                    q['point'] = 0.40 if q['part_type'] == 'PART_I' else 1.60
        elif preset == 'BGD_2025':
            if 'part1_total_points' not in exam_metadata:
                exam_metadata['part1_total_points'] = 6.0
            if 'part2_total_points' not in exam_metadata:
                exam_metadata['part2_total_points'] = 4.0
            if 'total_points' not in exam_metadata:
                exam_metadata['total_points'] = 10.0
            # Apply default points if question didn't have explicit point tag
            for q in parsed_questions:
                if not q.get('has_explicit_point'):
                    q['point'] = 0.25 if q['part_type'] == 'PART_I' else 1.00

        validation_summary = ExamValidator.validate_exam(parsed_questions, exam_metadata)

        return {
            'success': True,
            'exam_metadata': exam_metadata,
            'questions': parsed_questions,
            'summary': {
                'total_questions': len(parsed_questions),
                'part1_count': len(p1_questions),
                'part2_common_count': len(p2_common_questions),
                'part2_cs_count': len(p2_cs_questions),
                'part2_ict_count': len(p2_ict_questions),
                'matrix_preset': preset,
            },
            'validation_summary': validation_summary,
            'warnings': warnings
        }

    @classmethod
    def _parse_azota_footer_keys(cls, footer_text: str) -> Tuple[Dict[int, str], Dict[int, Dict[str, bool]]]:
        """
        Parses Azota footer answer format:
        1A 2B 3C or 1.A 2.B 3.C
        Câu 4: a)Đ b)S c)S d)Đ
        Câu 5: a)Đ b)S c)S d)S
        """
        p1_keys = {}
        p2_keys = {}
        if not footer_text:
            return p1_keys, p2_keys

        # Part 1 keys: e.g. 1A 2B 3C or 1.A 2.B
        p1_matches = re.finditer(r'(\d+)\s*[\.\:\-]?\s*([A-D])\b', footer_text, re.IGNORECASE)
        for m in p1_matches:
            q_num = int(m.group(1))
            ans_char = m.group(2).upper()
            p1_keys[q_num] = ans_char

        # Part 2 keys: e.g. Câu 4: a)Đ b)S c)S d)Đ or Câu 4: a)Dung b)Sai
        p2_lines = footer_text.split('\n')
        for line in p2_lines:
            q_match = re.search(r'Câu\s*(\d+)\s*:', line, re.IGNORECASE)
            if q_match:
                q_num = int(q_match.group(1))
                sub_matches = re.finditer(r'([a-d])\s*[\)\.\:]\s*(Đ|S|Đúng|Sai|Dung|True|False)', line, re.IGNORECASE)
                sub_dict = {}
                for sm in sub_matches:
                    lbl = sm.group(1).lower()
                    val_str = sm.group(2).lower()
                    is_true = val_str in ['đ', 'đúng', 'dung', 'true']
                    sub_dict[lbl] = is_true
                if sub_dict:
                    p2_keys[q_num] = sub_dict

        return p1_keys, p2_keys

    @classmethod
    def _parse_section_questions(
        cls,
        lines: List[str],
        part_type: str,
        branch: str,
        start_order_index: int,
        warnings: List[str]
    ) -> List[Dict[str, Any]]:
        questions = []
        raw_blocks: List[List[str]] = []
        current_block: List[str] = []

        for line in lines:
            if re.match(r'^\s*(\[CAU\s*\d*\]|[Cc][âa]u\s+\d+[:\.]|\[CAU\])', line, re.IGNORECASE):
                if current_block:
                    raw_blocks.append(current_block)
                current_block = [line]
            else:
                if current_block:
                    current_block.append(line)

        if current_block:
            raw_blocks.append(current_block)

        for q_idx, block in enumerate(raw_blocks):
            order_index = start_order_index + q_idx
            q_data = cls._parse_single_question_block(block, part_type, branch, order_index, warnings)
            if q_data:
                questions.append(q_data)

        return questions

    @classmethod
    def _parse_single_question_block(
        cls,
        block: List[str],
        part_type: str,
        branch: str,
        order_index: int,
        warnings: List[str]
    ) -> Dict[str, Any]:
        full_block_text = "\n".join(block)

        # Metadata tags: Competency, Difficulty & Point
        competency_category = 'PROG_BASIC'
        difficulty_level = 'TH'
        point = 0.50 if part_type == 'PART_I' else 2.00

        # Check for explicit branch tag in question block (e.g. [CS], [ICT], [CHUNG], [COMMON])
        if part_type == 'PART_II':
            if re.search(r'\[CS\]|\[NHANH_CS\]', full_block_text, re.IGNORECASE):
                branch = 'CS'
            elif re.search(r'\[ICT\]|\[NHANH_ICT\]', full_block_text, re.IGNORECASE):
                branch = 'ICT'
            elif re.search(r'\[CHUNG\]|\[COMMON\]|\[PHAN_CHUNG\]', full_block_text, re.IGNORECASE):
                branch = 'COMMON'

        # Check for point tag: [DIEM 0.5], [DIEM: 0.5], [0.5đ], (0.5đ), [0.5d], [0.5 diem], [0.5 điểm], [2.0đ], (2đ), etc.
        has_explicit_point = False
        point_match = re.search(
            r'\[(?:DIEM|SCORE|DIEM_SO)[\:\s]+([\d\.]+)\]|\((?:DIEM|SCORE|DIEM_SO)[\:\s]+([\d\.]+)\)|\[([\d\.]+)\s*(?:đ|d|diem|điểm|pts?|points?)\]|\(([\d\.]+)\s*(?:đ|d|diem|điểm|pts?|points?)\)',
            full_block_text,
            re.IGNORECASE
        )
        if point_match:
            pt_val = point_match.group(1) or point_match.group(2) or point_match.group(3) or point_match.group(4)
            try:
                point = float(pt_val)
                has_explicit_point = True
            except Exception:
                pass

        for cat_key in ['PROG_BASIC', 'ALGO_DS', 'OPTIMIZATION', 'DB_NETWORK', 'ICT_APP']:
            if f"[{cat_key}]" in full_block_text:
                competency_category = cat_key
                break

        for diff_key in ['NB', 'TH', 'VD', 'VDC']:
            if f"[{diff_key}]" in full_block_text or f",{diff_key}]" in full_block_text:
                difficulty_level = diff_key
                break

        # 1. Extract explanation if present
        explanation = ""
        exp_match = re.search(r'(\[HUONG_DAN\]|\[GIAI_THICH\]|\[LOI_GIAI\]|Hướng dẫn giải[:\.]|Lời giải[:\.])\s*([\s\S]*)', full_block_text, re.IGNORECASE)
        if exp_match:
            explanation = exp_match.group(2).strip()
            main_block_text = full_block_text[:exp_match.start()].strip()
        else:
            main_block_text = full_block_text.strip()

        # 2. Split into Prompt vs Options
        options = []
        raw_prompt_text = ""

        if part_type == 'PART_I':
            # AZOTA PART I: Scan all options *A. B. C. D., A*. A*) *A) at line start or separated by 2+ spaces / tabs
            # Exclude code dot-access like A.id, object.field, etc.
            opt_regex = re.compile(
                r'(?:^|(?<=\n)|(?<=\s{2})|(?<=\t))\s*(\*?[A-D]\*?[\.\:\)]|\*?[A-D]\*|\*[A-D])(?=\s|$)',
                re.MULTILINE | re.IGNORECASE
            )

            matches = list(opt_regex.finditer(main_block_text))
            if matches:
                raw_prompt_text = main_block_text[:matches[0].start()].strip()

                for i, m in enumerate(matches):
                    token = m.group(1).strip()
                    is_star = '*' in token
                    label = re.sub(r'[\*\.\:\)]', '', token).upper()

                    start_pos = m.end()
                    end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(main_block_text)
                    content = main_block_text[start_pos:end_pos].strip()

                    # Check for explicit "(Đáp án đúng)" or "[DUNG]"
                    if re.search(r'\(Đáp án đúng\)|\[ĐÚNG\]|\[DUNG\]|\(Đúng\)', content, re.IGNORECASE):
                        is_star = True
                        content = re.sub(r'\(Đáp án đúng\)|\[ĐÚNG\]|\[DUNG\]|\(Đúng\)', '', content, flags=re.IGNORECASE).strip()

                    options.append({
                        'label': label,
                        'content': content,
                        'is_correct': is_star,
                        'order_index': len(options) + 1,
                        'explanation': ''
                    })
            else:
                raw_prompt_text = main_block_text.strip()

            has_correct = any(opt['is_correct'] for opt in options)
            if not has_correct and options:
                warnings.append(f"Câu {order_index} (Phần I) chưa đánh dấu đáp án đúng (* hoặc màu đỏ).")
                # Do NOT auto-set options[0] as correct. Let teacher explicitly choose or leave unassigned.

        else:
            # AZOTA PART II: Scan *a)[0,NB] or a)[1,TH] or *a) or a*) or a.*
            sub_regex = re.compile(
                r'(?:^|(?<=\n)|(?<=\s{2})|(?<=\t))\s*(\*?[a-d]\*?[\.\:\)]|\*?[a-d]\*|\*[a-d])\s*(?:\[\d*\s*,\s*([A-Z]+)\]|\[([A-Z]+)\])?(?=\s|$)',
                re.MULTILINE | re.IGNORECASE
            )

            matches = list(sub_regex.finditer(main_block_text))
            if matches:
                raw_prompt_text = main_block_text[:matches[0].start()].strip()

                for i, m in enumerate(matches):
                    token = m.group(1).strip()
                    is_star = '*' in token
                    label = re.sub(r'[\*\.\:\)]', '', token).lower()
                    sub_diff = m.group(2) or m.group(3) or difficulty_level

                    start_pos = m.end()
                    end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(main_block_text)
                    content = main_block_text[start_pos:end_pos].strip()

                    # Check for explicit [DUNG] / [SAI]
                    if re.search(r'\[ĐÚNG\]|\[DUNG\]|\(Đúng\)|-\s*Đúng', content, re.IGNORECASE):
                        is_star = True
                        content = re.sub(r'\[ĐÚNG\]|\[DUNG\]|\(Đúng\)|-\s*Đúng', '', content, flags=re.IGNORECASE).strip()
                    elif re.search(r'\[SAI\]|\(Sai\)|-\s*Sai', content, re.IGNORECASE):
                        is_star = False
                        content = re.sub(r'\[SAI\]|\(Sai\)|-\s*Sai', '', content, flags=re.IGNORECASE).strip()

                    options.append({
                        'label': label,
                        'content': content,
                        'is_correct': is_star,
                        'difficulty_level': sub_diff,
                        'order_index': len(options) + 1,
                        'explanation': ''
                    })
            else:
                raw_prompt_text = main_block_text.strip()

            if len(options) < 4:
                warnings.append(f"Câu {order_index} Phần II ({branch}) có {len(options)}/4 ý.")

        # 3. Extract Code Snippets ONLY from raw_prompt_text
        code_snippets_list = []
        code_pattern = re.compile(r'```([a-zA-Z0-9_\+]*)\s*\n([\s\S]*?)```')
        for match in code_pattern.finditer(raw_prompt_text):
            lang = match.group(1).strip().lower() or 'python'
            if lang in ['c++', 'cplusplus']: lang = 'cpp'
            if lang in ['html5', 'htm']: lang = 'html'
            if lang in ['css3']: lang = 'css'
            snip_code = match.group(2).rstrip()
            snip_code = SmartCodeFormatter.format_code(snip_code, lang)
            code_snippets_list.append({
                'language': lang,
                'code': snip_code
            })

        tag_code_pattern = re.compile(r'\[CODE(?:_|\s+)?([a-zA-Z0-9_\+]*)\]([\s\S]*?)\[/CODE(?:_[a-zA-Z0-9_\+]*)?\]', re.IGNORECASE)
        for match in tag_code_pattern.finditer(raw_prompt_text):
            lang = match.group(1).strip().lower() or 'python'
            if lang in ['c++', 'cplusplus']: lang = 'cpp'
            if lang in ['html5', 'htm']: lang = 'html'
            if lang in ['css3']: lang = 'css'
            snip_code = match.group(2).rstrip()
            snip_code = SmartCodeFormatter.format_code(snip_code, lang)
            if not any(s['code'] == snip_code for s in code_snippets_list):
                code_snippets_list.append({
                    'language': lang,
                    'code': snip_code
                })

        # Remove code blocks ONLY from raw_prompt_text
        cleaned_prompt_text = code_pattern.sub('', raw_prompt_text)
        cleaned_prompt_text = tag_code_pattern.sub('', cleaned_prompt_text)

        # Build combined code snippet or single code snippet
        if len(code_snippets_list) > 1:
            code_snippet = "\n\n".join([f"```{s['language']}\n{s['code']}\n```" for s in code_snippets_list])
            code_language = code_snippets_list[0]['language']
        elif len(code_snippets_list) == 1:
            code_snippet = code_snippets_list[0]['code']
            code_language = code_snippets_list[0]['language']
        else:
            code_snippet = ""
            code_language = "python"

        # Clean header tags like [CAU 1], [DIEM ...], [PROG_BASIC] from prompt text
        prompt_lines = [l for l in cleaned_prompt_text.split('\n') if l.strip()]
        if prompt_lines:
            first_line = prompt_lines[0]
            cleaned_first = re.sub(r'^\s*(\[CAU\s*\d*\]|[Cc][âa]u\s+\d+[:\.]|\[CAU\])', '', first_line, flags=re.IGNORECASE)
            cleaned_first = re.sub(r'\[(PROG_BASIC|ALGO_DS|OPTIMIZATION|DB_NETWORK|ICT_APP|NB|TH|VD|VDC|CS|ICT|CHUNG|COMMON|NHANH_CS|NHANH_ICT|PHAN_CHUNG)\]', '', cleaned_first, flags=re.IGNORECASE)
            cleaned_first = re.sub(r'\[(?:DIEM|SCORE|DIEM_SO)[\:\s]+[\d\.]+\]|\((?:DIEM|SCORE|DIEM_SO)[\:\s]+[\d\.]+\)|\[[\d\.]+\s*(?:đ|d|diem|điểm|pts?|points?)\]|\([\d\.]+\s*(?:đ|d|diem|điểm|pts?|points?)\)', '', cleaned_first, flags=re.IGNORECASE)
            prompt_lines[0] = cleaned_first

        # AI Heuristic Auto-Detection if no explicit code fences were used
        if not code_snippet:
            clean_p_lines, auto_code, auto_lang = SmartCodeDetector.extract_unformatted_code(prompt_lines)
            if auto_code:
                code_snippet = auto_code
                code_language = auto_lang
                prompt_lines = clean_p_lines

        content_final = "\n".join([p.strip() for p in prompt_lines if p.strip()]).strip()

        return {
            'order_index': order_index,
            'part_type': part_type,
            'branch': branch,
            'point': point,
            'has_explicit_point': has_explicit_point,
            'content': content_final or f"Câu hỏi {order_index}",
            'code_snippet': code_snippet,
            'code_language': code_language,
            'competency_category': competency_category,
            'difficulty_level': difficulty_level,
            'options': options,
            'explanation': explanation
        }


class ExamValidator:
    """
    Comprehensive Exam & Question Quality Diagnostics Engine.
    Detects missing options, missing correct answers, duplicate labels, empty prompts,
    invalid formatting, and matrix mismatches.
    """

    @classmethod
    def validate_question(cls, q: Dict[str, Any], global_index: int = 0) -> Dict[str, Any]:
        """
        Validates a single question and populates its 'issues', 'has_error', 'has_warning'.
        """
        issues: List[Dict[str, str]] = []
        part_type = q.get('part_type', 'PART_I')
        order_index = q.get('order_index', global_index + 1)
        branch = q.get('branch', 'COMMON')
        content = (q.get('content') or '').strip()
        code_snippet = (q.get('code_snippet') or '').strip()
        options = q.get('options') or []
        explanation = (q.get('explanation') or '').strip()

        # 1. Check Empty Prompt
        if not content and not code_snippet:
            issues.append({
                'type': 'error',
                'code': 'EMPTY_PROMPT',
                'message': 'Nội dung câu hỏi đang bị rỗng.',
                'suggestion': 'Nhập nội dung câu hỏi hoặc mã nguồn cho câu này.'
            })
        elif len(content) < 3 and not code_snippet:
            issues.append({
                'type': 'warning',
                'code': 'SHORT_PROMPT',
                'message': f"Nội dung câu hỏi quá ngắn ('{content}').",
                'suggestion': 'Kiểm tra lại xem câu hỏi có bị thiếu chữ hay không.'
            })

        # 2. Check Options Count
        if not options:
            issues.append({
                'type': 'error',
                'code': 'NO_OPTIONS',
                'message': 'Câu hỏi không có bất kỳ phương án lựa chọn nào.',
                'suggestion': 'Thêm các phương án A, B, C, D (Phần I) hoặc a, b, c, d (Phần II).'
            })
        elif part_type == 'PART_I':
            if len(options) < 4:
                existing_labels = ", ".join([opt.get('label', '?') for opt in options])
                issues.append({
                    'type': 'error',
                    'code': 'PART_I_INCOMPLETE_OPTIONS',
                    'message': f"Phần I chỉ có {len(options)}/4 phương án (Hiện có: {existing_labels}).",
                    'suggestion': 'Bổ sung thêm các phương án còn thiếu để đủ 4 phương án A, B, C, D.'
                })
        elif part_type == 'PART_II':
            if len(options) < 4:
                existing_labels = ", ".join([opt.get('label', '?') for opt in options])
                issues.append({
                    'type': 'error',
                    'code': 'PART_II_INCOMPLETE_SUBITEMS',
                    'message': f"Phần II chỉ có {len(options)}/4 ý (Hiện có: {existing_labels}).",
                    'suggestion': 'Bổ sung các ý con còn thiếu để đủ 4 ý a), b), c), d).'
                })

        # 3. Check Option Content and Labels
        labels_seen = set()
        for idx, opt in enumerate(options):
            label = (opt.get('label') or '').strip()
            opt_content = (opt.get('content') or '').strip()

            if not label:
                issues.append({
                    'type': 'error',
                    'code': 'MISSING_OPTION_LABEL',
                    'message': f"Phương án thứ {idx + 1} chưa có nhãn (A/B/C/D hoặc a/b/c/d).",
                    'suggestion': 'Đặt nhãn chuẩn cho phương án.'
                })
            elif label in labels_seen:
                issues.append({
                    'type': 'error',
                    'code': 'DUPLICATE_OPTION_LABELS',
                    'message': f"Trùng lặp nhãn phương án '{label}' trong cùng câu hỏi.",
                    'suggestion': f"Đổi tên nhãn phương án '{label}' để không bị trùng."
                })
            labels_seen.add(label)

            if not opt_content:
                issues.append({
                    'type': 'error',
                    'code': 'EMPTY_OPTION_CONTENT',
                    'message': f"Nội dung của phương án {label} đang bị rỗng.",
                    'suggestion': f"Nhập nội dung cho phương án {label}."
                })

        # 4. Check Multiple Correct Answers in Part I
        if options and part_type == 'PART_I':
            correct_count = sum(1 for opt in options if opt.get('is_correct'))
            if correct_count > 1:
                issues.append({
                    'type': 'error',
                    'code': 'PART_I_MULTIPLE_CORRECT',
                    'message': f"Phần I có tới {correct_count} đáp án được đánh dấu đúng.",
                    'suggestion': 'Chỉ chọn duy nhất 1 đáp án đúng cho câu hỏi Phần I.'
                })

        has_error = any(issue['type'] == 'error' for issue in issues)
        has_warning = any(issue['type'] == 'warning' for issue in issues)

        q['issues'] = issues
        q['has_error'] = has_error
        q['has_warning'] = has_warning
        return q

    @classmethod
    def validate_exam(cls, questions: List[Dict[str, Any]], metadata: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """
        Validates all questions in the exam and generates a complete diagnostic summary.
        """
        total = len(questions)
        error_indices: List[int] = []
        warning_indices: List[int] = []
        detailed_issues: List[Dict[str, Any]] = []

        for idx, q in enumerate(questions):
            cls.validate_question(q, global_index=idx)
            if q.get('has_error'):
                error_indices.append(idx)
            if q.get('has_warning'):
                warning_indices.append(idx)

            if q.get('issues'):
                detailed_issues.append({
                    'list_index': idx,
                    'order_index': q.get('order_index', idx + 1),
                    'part_type': q.get('part_type', 'PART_I'),
                    'branch': q.get('branch', 'COMMON'),
                    'has_error': q.get('has_error', False),
                    'has_warning': q.get('has_warning', False),
                    'issues': q['issues']
                })

        error_count = len(error_indices)
        warning_count = len(warning_indices)
        valid_count = total - error_count

        return {
            'total_questions': total,
            'valid_count': valid_count,
            'error_count': error_count,
            'warning_count': warning_count,
            'is_perfect': error_count == 0 and warning_count == 0 and total > 0,
            'is_valid': error_count == 0 and total > 0,
            'error_question_indices': error_indices,
            'warning_question_indices': warning_indices,
            'detailed_issues': detailed_issues
        }

